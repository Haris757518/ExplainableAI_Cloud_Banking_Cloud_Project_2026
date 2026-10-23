"""Build the research submission from verified final local results only.

Execute with the bundled document Python after its operation marker.
"""
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def main():
    final = json.loads((ROOT/'results/final_validation/report.json').read_text())
    cleanup = json.loads((ROOT/'results/final_validation/aws-cleanup.json').read_text())
    assert final['status'] == 'evaluated' and cleanup['status'] == 'terminated' and cleanup['root_volume_deleted']
    assert final['export_sha256'] == cleanup['final_export_sha256']
    initial = json.loads((ROOT/'results/banking_study/v1/report.json').read_text())
    calibration = json.loads((ROOT/'results/probe_xai/calibrated-detection.json').read_text())
    graph = final['method_summary']['graph']
    seed42 = next(s for s in graph['seed_results'] if s['seed'] == 42)
    n = graph['independent_executed_fault_repetitions']
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Header','Footer']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(9)
    normal.paragraph_format.line_spacing=1.12
    doc.styles['Title'].font.size=Pt(26)
    doc.styles['Heading 1'].font.size=Pt(19)
    doc.styles['Heading 2'].font.size=Pt(13)
    header=section.header.paragraphs[0];header.text='Haris  •  Banking investigation research  •  October 2026'
    header.runs[0].font.size=Pt(9)
    footer=section.footer.paragraphs[0];footer.text='Research prototype  |  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

    def p(text): doc.add_paragraph(text)
    def h(text): doc.add_heading(text,level=2)
    def page(title): doc.add_page_break();doc.add_heading(title,level=1)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for c,x in zip(t.rows[0].cells,headers):
            c.text=x
            for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF1');c._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            for c,x in zip(t.add_row().cells,row):c.text=str(x)
        for row in t.rows:
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        doc.add_paragraph()

    doc.add_paragraph('Explainable AI for cloud banking investigation',style='Title')
    doc.add_paragraph('Executed workload validation and corrective code study',style='Subtitle')
    p('This report explains the implemented banking investigation prototype, the evidence supporting its findings, and the limits of its research claims. It is written for project examiners and readers who need to understand what failed, why it failed, what changed and how recovery was checked.')
    h('Main findings')
    p('The application connects anomaly detection with service investigation and transaction-integrity checks. A controlled test exposed a reader defect: transaction numbers could advance before an earlier transfer committed, leaving displayed balances or history incomplete. The corrected readers query committed ledger records directly. They passed three before-and-after AWS trials and reader restart checks, with deployed build checksums verified.')
    p(f"In the new frozen-rule AWS workload, evidence-guided localization identified the expected service at {seed42['refined_top1_count']} of {n} CPU and database fault checkpoints using seed 42. GNN ranking alone identified {seed42['gnn_top1_count']}. Results for every seed and both graph variants appear later in this report. These controlled tests do not establish production accuracy or graph superiority.")
    h('Research contribution')
    p('The proposed PROBE-XAI contribution is an integrated investigation workflow: inspect direct evidence progressively, keep model symptom rankings separate from cause attribution, explain the deepest supported mechanism, and require ledger integrity before declaring recovery. The demonstrated integrity gate rejected an apparently fast but incorrect account state. Worldwide novelty remains unproven.')
    h('Completion boundary')
    p('The software, controlled cloud experiment, verified evidence export and exact-instance cleanup are implemented. The deliverable is a completed research prototype for an isolated demonstration. The targeted prior-work and production-readiness assessments later in this report explain why a worldwide-first claim and live customer banking deployment are not supported.')

    page('Data and executed environment')
    p('The study keeps three evidence families separate. RCAEval RE2 supplies controlled microservice recordings for benchmark research. Bank of Anthos supplies actual banking requests, resource measurements and PostgreSQL records from an isolated AWS workload. Published Monzo and Starling aggregate statistics provide banking context, without serving as internal service fault labels.')
    table(['Evidence family','Use','Boundary'],[
        ['RCAEval RE2','15 recorded cases and per-recording models','Retrospective benchmark audit'],
        ['Original AWS banking run','3,025 submitted requests and ledger matches','Acknowledgement parser defect excluded'],
        ['New AWS banking run',f"{final['submitted_requests']:,} submissions; {final['ledger_matches']:,} ledger matches",'One host and controlled repetitions'],
        ['Published bank history','Independent aggregate reference','No joins to fault labels'],
    ])
    h('Cloud execution')
    p('Bank of Anthos ran on one disposable Ubuntu EC2 host using single-node K3s. Banking services remained private. The run reused the authorized lab profile, used a bounded runtime and retained an independent guest shutdown deadline. It created no EKS cluster, managed database or load balancer. Only the exact tagged research instance was terminated; the pre-existing payment server was preserved.')
    p('Measurements include request starts and completions, frontend client spans, actual cgroup processor limits and usage, memory readings, database lock waits and ledger rows. Missing backend or SQL trace coverage remains missing. The observed frontend graph does not invent database trace edges.')
    h('Evidence retention')
    p('The final archive was downloaded and checksum verified before cleanup. Per-file hashes bind the measurements, collector and ledger snapshot. The database dump was checked before backup. Cleanup confirmed instance termination and root-volume deletion. A prior repeat completed but missed its raw export before the deadline; its summary remains documented and is excluded from localization scoring.')

    page('Detection and investigation method')
    h('Chronological model study')
    p('The banking graph autoencoder predicts the next measured feature bin for four services. Each bin preserves missing values with masks. The normal prefix provides 216 seconds for training, the next 72 seconds for validation and a further 72 seconds for normal holdout. Scaling and checkpoint selection use the permitted prefix. Graph and identity variants share training settings across seeds 42, 43 and 44. Persistence and training-mean baselines remain available in the model study.')
    p('A GNN residual indicates an unexpected measurement. A large residual can appear in a downstream service, so the residual ranking is retained as a symptom ranking. It does not itself establish the originating service.')
    h('Progressive direct evidence')
    p('The investigator checks whether processor allowance fell while demand stayed near its reference, whether database lock waits and blocked sessions explain waiting, or whether increased incoming traffic accompanies slower responses. It preserves multiple supported contributors and abstains when instruments or the inspection budget cannot distinguish them. Field-inspection units are an in-memory effort measure, not a measured AWS query bill or information-gain guarantee.')
    p('The original banking GNN ranked the wrong service at both inspected fault checkpoints. Adding direct evidence corrected those two development cases. The later evaluation freezes rules before a new workload, writes predictions before opening intervention labels, and scores fixed checkpoints 30 seconds after each CPU or database fault begins.')
    h('Benchmark inspection effort')
    p('In the retained 15-case RE2 comparison, progressive inspection examined 2.0 follow-up columns on average, compared with 4.8 for conventional inspection. Both retained 60.0 percent root-service Top 1 and 13.3 percent mechanism agreement; both abstained from an underlying-driver claim in every case. The statistical service baseline reached 73.3 percent Top 1. These results support reduced column inspection in this bounded comparison, without establishing graph superiority, live query-cost savings or labelled deeper-driver accuracy.')
    h('Recovery requires correctness')
    p('Recovery requires both useful response behavior and transaction checks. A ledger match is distinct from a caller acknowledgement. The full workflow compares both account balances with committed ledger records, checks unique transaction history, verifies later transfers and checks successful reads without continuing database blocking.')

    page('Measured results')
    table(['Variant','Seed','GNN top one','Evidence top one','Fault detected'],[
        [method,str(s['seed']),f"{s['gnn_top1_count']}/{n}",f"{s['refined_top1_count']}/{n}",f"{s['detected_count']}/{n}"]
        for method in ['graph','identity'] for s in final['method_summary'][method]['seed_results']])
    p('Each denominator counts executed CPU or database fault repetitions in the same workload. Random seeds are model repetitions, not additional independent faults. The graph and identity variants use the same direct-evidence investigator. Any localization benefit from that investigator does not establish a benefit from graph structure.')
    normal=[r for r in final['model_study']['methods'] if r['method']=='graph']
    p('Normal holdout graph alerts by seed: '+', '.join(f"seed {r['seed']}: {r['normal_holdout_alerts']} of {r['normal_holdout_windows']} observed bins" for r in normal)+'. These short-holdout counts are reported separately from fault detection and do not establish operational reliability.')
    table(['New workload accounting','Observed count'],[
        ['Submitted requests',final['submitted_requests']],['Unique matching ledger records',final['ledger_matches']],
        ['Acknowledged successes',final['acknowledged_success']],['Uncertain caller outcomes',final['uncertain_requests']],
        ['Other responses not classified as success',final['submitted_requests']-final['acknowledged_success']-final['uncertain_requests']],
    ])
    h('RE2 calibration audit')
    p(f"Across 15 already inspected recordings, calibrated normal-checkpoint alerts fell from 11 to {calibration['confusion']['false_positive']}. The detector found {calibration['confusion']['true_positive']} of 15 post-injection checkpoints and missed {calibration['confusion']['false_negative']}. Precision was {calibration['precision']:.3f}, recall {calibration['recall']:.3f} and F1 {calibration['f1']:.3f}. Calibration uses a trailing three-point residual and a validation-prefix maximum multiplied by 1.5. This retrospective checkpoint audit is not an independent production false-alarm study.")

    page('Reader defect correction and recovery gate')
    h('Why the defect happens')
    p('A transaction number is allocated before its transaction necessarily commits. If transaction 74 commits before transaction 73, a reader that remembers only the highest number can move its cursor past 73. When 73 later commits, that reader can skip it. The controlled experiment held the earlier transaction open and committed the later one first, reproducing a 17-cent discrepancy and a missing history entry.')
    h('Implemented correction')
    p('Authorized balance requests now calculate the account balance from committed ledger records. Authorized history requests query the committed records in deterministic timestamp and transaction-number order. Database errors return an error rather than a stale cached result. The background cache may remain present, but it is no longer authoritative for these endpoints.')
    table(['Check','Result'],[
        ['Original code with reversed commit order','3 of 3 trials incorrect'],
        ['Corrected readers with reversed commit order','3 of 3 trials correct'],
        ['Reader restart','Both balances match the ledger; diagnostic transfers occur once in both histories, without duplicate IDs'],
        ['Local Java regression suite','19 tests passed'],
        ['Deployed class checksums','Match the validated build'],
    ])
    p('These tests cover the reproduced defect and the tested restart behavior. Direct ledger reads increase database work. They do not certify every transaction isolation pattern, production scale or failure mode.')
    h('Why speed alone was insufficient')
    p('Before reader resynchronization in the original workload, three balance reads completed below the demonstrative 100 ms policy while the balance differed from the ledger by 125,678 cents. A speed-only recovery policy passed. The integrity gate rejected that state. After reconciliation, the discrepancy was zero and both policies passed. These are two observation sets from one workload, not independent recovery trials.')

    page('Use reproduction and research scope')
    h('End user demonstration')
    p('Start the project launcher and choose Run banking investigation. Select the balance finding to see the missed transfer, affected reader, commit-order explanation, implemented correction and verification result. The processor, database and traffic findings explain only what their instruments support. Research and Validation contains the model comparisons and their limits. The normal screen presents readable explanations rather than raw JSON.')
    h('Local reproduction')
    p('Run scripts/replay_banking_investigation.py from the project directory to verify the bundled original recording and later correction proof, then produce a readable investigation report without AWS. The final study is reproduced with scripts/evaluate_final_banking_validation.py and the checksum-verified export directory. The prospective protocol, amendments, predictions written before label scoring, model checkpoints and report are retained under results/final_validation.')
    p('Verification completed during this update includes 195 Python tests, 19 Java tests, a frontend production build and four browser acceptance checks. Credentials and private keys remain excluded from the repository. The final archive and SQL backup remain on the D drive.')
    h('Related work and claim limits')
    p('The review below distinguishes established research ideas from the project-specific engineering and executed case study. The contribution is evidence-limited banking investigation with an integrity gate, illustrated by a reproduced commit-order defect and a verified correction. It is not a claim of a new general-purpose GNN algorithm.')

    page('Novelty assessment')
    p('Review date: 23 October 2026. A targeted English-language primary-source review examined anomaly localization, progressive diagnosis, graph explanations and recovery-aware investigation. The comparison includes conference papers and clearly identified 2026 preprints. It is not an exhaustive worldwide literature, patent or commercial-product search; absence of an exact search match would not establish priority.')
    table(['Prior work','Established overlap with this project'],[
        ['Eadro, ICSE 2023 [1]','Joint anomaly detection and root-cause localization using multiple observability sources.'],
        ['MicroRCA, NOMS 2020 [3]','Graph-based localization connecting application symptoms and resource evidence.'],
        ['GNNExplainer, 2019 [4]','Explaining a graph model through graph structure and feature importance.'],
        ['PDA, LISA 2007 [8]','Health checks and scoped diagnostic probes that drill down progressively.'],
        ['R2Act, July 2026 preprint [6]','Separating diagnosis from the validity of recovery plans and actions.'],
        ['Trajectory-level RCA study, v2, October 2026 [7]','Studying diagnostic evidence and the investigation process beyond fault localization.'],
    ])
    h('What this project demonstrates')
    p('The specific contribution is a transparent banking case study that separates model symptoms from direct resource evidence, limits deeper-cause claims to observed mechanisms, reproduces a commit-order reader defect, and rejects fast but financially incorrect recovery. The retained code, before-and-after trials and ledger checks make that engineering result reviewable.')
    h('Claim permitted by the evidence')
    p('An implemented integration and demonstrated case-study contribution are supported. Worldwide-first novelty, general algorithmic superiority and independently validated deeper-driver accuracy are not established. The comparison identifies substantial prior overlap; a unique combination alone is not proof of research priority. Future priority claims would require a systematic search and a precise distinction from the closest methods.')

    page('Production-readiness assessment')
    p('Assessment date: 23 October 2026. This is an evidence-based review of the inspected release and retained experiments, not a security certification. AWS reliability guidance calls for validation of functional and nonfunctional requirements, including scalability, performance and resilience [9]; identity and access controls are separate requirements [10].')
    table(['Area','Observed evidence and remaining gate'],[
        ['Functional behavior','195 Python tests, 19 Java tests, a frontend build and four browser checks passed. These support the tested research workflow.'],
        ['Access control','Inspected dashboard API routes have no authentication or role enforcement. Banking JWT validation does not protect these investigation routes. Add and verify operator access controls before wider deployment.'],
        ['Availability','The executed environment used one EC2 host and single-node K3s. Multi-host redundancy and failover were not demonstrated.'],
        ['Backup and recovery','An exported SQL backup was checksum verified. A complete restore and recovery-time test was not demonstrated.'],
        ['Capacity and service targets','Only bounded controlled traffic was evaluated. Agree service targets and test sustained load, overload and the extra database cost of authoritative ledger reads.'],
        ['Diagnostic coverage','Six controlled faults in one workload and a short normal holdout do not establish operational accuracy. Backend SQL traces and queue visibility remain incomplete.'],
        ['Security and operations','No independent penetration-test or production incident-response exercise is present in the inspected evidence. Verify hardening, secrets lifecycle, monitoring and release rollback.'],
    ])
    h('Decision')
    p('Ready for an isolated, supervised research demonstration using the recorded evidence. Not ready for real customer banking on the inspected evidence. The research implementation can be complete while production engineering remains a separate undertaking. No new deployment or paid resources are required to reproduce the retained demonstration.')
    page('Sources and retained evidence')
    p('The references below establish the related-work context and the source application. Project measurements come from executed workloads and retained verification files, rather than from these publications.')
    for text in [
        '[1] Eadro. ICSE 2023. https://arxiv.org/abs/2302.05092',
        '[2] RCAEval. WWW Companion 2025. https://arxiv.org/abs/2412.17015',
        '[3] MicroRCA. NOMS 2020. https://doi.org/10.1109/NOMS47738.2020.9110353',
        '[4] GNNExplainer. NeurIPS 2019. https://arxiv.org/abs/1903.03894',
        '[5] Bank of Anthos source. https://github.com/GoogleCloudPlatform/bank-of-anthos',
        '[6] R2Act: Can LLMs Really Recover Microservice Failures? A Recovery-Aware Evaluation of Diagnosis-to-Action Reasoning. July 2026 preprint. https://arxiv.org/abs/2607.04623',
        '[7] Beyond Fault Localization: A Trajectory-Level Study of LLM Agents for Microservice Root Cause Analysis. Preprint v2, 9 October 2026. https://arxiv.org/abs/2608.21310',
        '[8] PDA: A Tool for Automated Problem Determination. LISA 2007. https://www.usenix.org/conference/lisa-07/presentation/pda-tool-automated-problem-determination',
        '[9] AWS Well-Architected, Test reliability. https://docs.aws.amazon.com/wellarchitected/latest/reliability-pillar/test-reliability.html',
        '[10] AWS Well-Architected, Identity and access management. https://docs.aws.amazon.com/wellarchitected/latest/security-pillar/identity-and-access-management.html',
    ]:p(text)
    page('Retained evidence and assessment traceability')
    h('Evidence index')
    table(['Project record','Location under the project directory'],[
        ['Original banking recording','results/aws_workload/replay/'],
        ['Reader correction and deployed build proof','results/correctness_fix/aws-proof/'],
        ['Frozen repeat protocol and scoring plan','results/final_validation/'],
        ['Portable repeat measurements','results/final_validation/raw/'],
        ['Final scores and verified cleanup','results/final_validation/report.json and aws-cleanup.json'],
        ['Retrospective RE2 calibration audit','results/probe_xai/calibrated-detection.json'],
    ])
    p('Full SQL backups remain on D:. The portable research bundle contains allowlisted measurements, ledger records, collector code and source checksums. Credentials and private keys are excluded. Reproduction instructions are in docs/BANKING_DEMO_GUIDE.md.')
    p('Readiness inspection covered backend/main.py, backend/cloud_investigation_api.py, deployment scripts and retained validation results. No restore, failover or independent security-test receipt was identified in this inspected release. This records the evidence boundary rather than asserting that all possible checks have failed.')
    # Remove inherited title rules so the report has a plain, black heading hierarchy.
    for element in [s.element for s in doc.styles] + [para._p for para in doc.paragraphs]:
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    output=ROOT/'submission/Banking_Investigation_Final_Assessed_Report_20261023.docx'
    assert not output.exists(), 'Use a new output filename for a new revision'
    doc.save(output)
    print(output)


if __name__=='__main__':main()
